#!/usr/bin/env python3
import openpyxl

# OpenPyXL supports creating bar, line, scatter, and pie charts using the data in a sheet’s cells. 
# To make a chart, you need to do the following:

wb = openpyxl.Workbook()
sheet = wb.active

# 0. create some sample data in column A or use existing
for i in range(1, 11): 
    sheet['A' + str(i)] = i

for i in range(1, 5): 
    sheet['B' + str(i)] = f'Test {i}'
    sheet['C' + str(i)] = (i * 10)

# 1. create a Reference object from a rectangular selection of cells
refObj = openpyxl.chart.Reference(sheet, min_col=1, min_row=1, max_col=1, max_row=10)

refObjPieData = openpyxl.chart.Reference(sheet, min_col=3, min_row=1, max_col=3, max_row=4)
refObjPieLabels = openpyxl.chart.Reference(sheet, min_col=2, min_row=1, max_col=2, max_row=4)

# 2. create a Series object by passing in the Reference object
seriesObj = openpyxl.chart.Series(refObj, title='Some series')
seriesForPie = openpyxl.chart.Series(refObjPieData, title='Pieee series')

# 3. create a Chart object
chartObj = openpyxl.chart.BarChart()
chartObj.title = 'My First Chart'
# You can also create line charts, scatter charts, and pie charts by calling:
#  openpyxl.chart.LineChart()
#  openpyxl.chart.ScatterChart()
#  openpyxl.chart.PieChart()
# respectively
pieObj = openpyxl.chart.PieChart()
pieObj.title = 'I like pies!'

# 4. append the Series object to the Chart object.
chartObj.append(seriesObj)
pieObj.append(seriesForPie)
pieObj.set_categories(refObjPieLabels)

# 5. add the chart object to the Worksheet object, optionally specifying which cell 
# the top left corner of the chart should be positioned..
sheet.add_chart(chartObj, 'C5')
sheet.add_chart(pieObj, 'C21')

wb.save('sampleChart.xlsx')